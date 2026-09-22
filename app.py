import os
import uuid
import threading
from pathlib import Path

import numpy as np
import faiss

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf

from docx import Document
from pptx import Presentation

from flask import Flask, render_template, request, jsonify
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

from google import genai
from google.genai import types

from sentence_transformers import SentenceTransformer


# =========================================================
# ENVIRONMENT
# =========================================================

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if GEMINI_API_KEY:
    client = genai.Client(api_key=GEMINI_API_KEY)
else:
    client = None


# =========================================================
# FLASK
# =========================================================

app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent

UPLOAD_FOLDER = BASE_DIR / "uploads"
UPLOAD_FOLDER.mkdir(parents=True, exist_ok=True)

app.config["UPLOAD_FOLDER"] = str(UPLOAD_FOLDER)


# =========================================================
# FILE SETTINGS
# =========================================================

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "txt",
    "pptx"
}


# =========================================================
# MODELS
# =========================================================

# Gemini is used ONLY for answer generation.
GENERATION_MODEL = "gemini-3.1-flash-lite"

# Local embedding model.
LOCAL_EMBEDDING_MODEL = "all-MiniLM-L6-v2"

# all-MiniLM-L6-v2 produces 384-dimensional vectors.
EMBEDDING_DIMENSION = 384


# =========================================================
# RAG SETTINGS
# =========================================================

CHUNK_SIZE = 1800
CHUNK_OVERLAP = 250

TOP_K = 5
SIMILARITY_THRESHOLD = 0.25

LOCAL_EMBEDDING_BATCH_SIZE = 64

MAX_RETRIES = 3


# =========================================================
# GLOBAL DATA
# =========================================================

documents = {}

chunks = []

faiss_index = None

rag_lock = threading.Lock()

embedding_model = None

embedding_model_lock = threading.Lock()


# =========================================================
# LOCAL EMBEDDING MODEL
# =========================================================

def get_embedding_model():

    global embedding_model

    if embedding_model is None:

        with embedding_model_lock:

            if embedding_model is None:

                print("\n====================================")
                print("Loading local embedding model...")
                print(f"Model: {LOCAL_EMBEDDING_MODEL}")
                print("====================================\n")

                embedding_model = SentenceTransformer(
                    LOCAL_EMBEDDING_MODEL
                )

                print("Local embedding model loaded successfully.\n")

    return embedding_model


# =========================================================
# FILE HELPERS
# =========================================================

def allowed_file(filename):

    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower()
        in ALLOWED_EXTENSIONS
    )


def get_extension(filename):

    return filename.rsplit(".", 1)[1].lower()


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(text):

    if not text:
        return ""

    lines = []

    for line in text.splitlines():

        line = " ".join(line.split())

        if line:
            lines.append(line)

    return "\n".join(lines).strip()


# =========================================================
# PDF EXTRACTION
# =========================================================

def extract_pdf(file_path):

    pages = []

    with pymupdf.open(file_path) as pdf:

        for page_number, page in enumerate(
            pdf,
            start=1
        ):

            text = page.get_text("text")

            text = clean_text(text)

            if text:

                pages.append({
                    "page": page_number,
                    "text": text
                })

    return pages


# =========================================================
# DOCX EXTRACTION
# =========================================================

def extract_docx(file_path):

    document = Document(file_path)

    pages = []

    section_number = 0

    for paragraph in document.paragraphs:

        text = clean_text(paragraph.text)

        if text:

            section_number += 1

            pages.append({
                "page": section_number,
                "text": text
            })

    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = clean_text(cell.text)

                if text:
                    cells.append(text)

            if cells:

                section_number += 1

                pages.append({
                    "page": section_number,
                    "text": " | ".join(cells)
                })

    return pages


# =========================================================
# TXT EXTRACTION
# =========================================================

def extract_txt(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        text = file.read()

    text = clean_text(text)

    if not text:
        return []

    return [{
        "page": 1,
        "text": text
    }]


# =========================================================
# PPTX EXTRACTION
# =========================================================

def extract_pptx(file_path):

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = clean_text(shape.text)

                if text:

                    slide_text.append(text)

        combined_text = "\n".join(
            slide_text
        ).strip()

        if combined_text:

            slides.append({
                "page": slide_number,
                "text": combined_text
            })

    return slides


# =========================================================
# DOCUMENT EXTRACTION ROUTER
# =========================================================

def extract_document(file_path, extension):

    if extension == "pdf":
        return extract_pdf(file_path)

    if extension == "docx":
        return extract_docx(file_path)

    if extension == "txt":
        return extract_txt(file_path)

    if extension == "pptx":
        return extract_pptx(file_path)

    return []


# =========================================================
# CHUNKING
# =========================================================

def create_chunks(
    pages,
    filename,
    document_id
):

    result = []

    for page_data in pages:

        page_number = page_data["page"]

        text = page_data["text"]

        if not text:
            continue

        start = 0

        text_length = len(text)

        while start < text_length:

            end = min(
                start + CHUNK_SIZE,
                text_length
            )

            chunk_text = text[start:end].strip()

            if chunk_text:

                result.append({

                    "id": uuid.uuid4().hex,

                    "document_id": document_id,

                    "filename": filename,

                    "page": page_number,

                    "text": chunk_text

                })

            if end >= text_length:
                break

            start = end - CHUNK_OVERLAP

    return result


# =========================================================
# LOCAL EMBEDDINGS
# =========================================================

def create_embeddings(
    texts,
    progress_callback=None
):

    if not texts:
        return []

    model = get_embedding_model()

    total = len(texts)

    all_vectors = []

    for start in range(
        0,
        total,
        LOCAL_EMBEDDING_BATCH_SIZE
    ):

        end = min(
            start + LOCAL_EMBEDDING_BATCH_SIZE,
            total
        )

        batch = texts[start:end]

        print(
            f"Embedding chunks "
            f"{start + 1}-{end} of {total}"
        )

        try:

            vectors = model.encode(

                batch,

                batch_size=LOCAL_EMBEDDING_BATCH_SIZE,

                show_progress_bar=False,

                convert_to_numpy=True,

                normalize_embeddings=True

            )

        except Exception as error:

            print("\nEmbedding error:")

            print(str(error))

            raise

        vectors = np.asarray(
            vectors,
            dtype=np.float32
        )

        all_vectors.extend(vectors)

        if progress_callback:

            progress_callback(
                end,
                total
            )

    return all_vectors


# =========================================================
# ADD DOCUMENT TO FAISS
# =========================================================

def add_document_to_index(
    new_chunks,
    progress_callback=None
):

    global faiss_index

    if not new_chunks:
        return 0

    texts = [
        item["text"]
        for item in new_chunks
    ]

    vectors = create_embeddings(
        texts,
        progress_callback
    )

    if len(vectors) != len(new_chunks):

        raise RuntimeError(
            "Embedding count does not match chunk count."
        )

    matrix = np.asarray(
        vectors,
        dtype=np.float32
    )

    with rag_lock:

        if faiss_index is None:

            faiss_index = faiss.IndexFlatIP(
                EMBEDDING_DIMENSION
            )

        faiss_index.add(matrix)

        chunks.extend(new_chunks)

    return len(new_chunks)


# =========================================================
# VECTOR SEARCH
# =========================================================

def search_documents(question):

    with rag_lock:

        if faiss_index is None:
            return []

        if len(chunks) == 0:
            return []

        total_chunks = len(chunks)

        limit = min(
            TOP_K,
            total_chunks
        )

    query_vectors = create_embeddings(
        [question]
    )

    if not query_vectors:
        return []

    query_vector = np.asarray(
        query_vectors[0],
        dtype=np.float32
    ).reshape(1, -1)

    with rag_lock:

        scores, positions = faiss_index.search(
            query_vector,
            limit
        )

        results = []

        for score, position in zip(
            scores[0],
            positions[0]
        ):

            if position < 0:
                continue

            score = float(score)

            if score < SIMILARITY_THRESHOLD:
                continue

            item = chunks[int(position)].copy()

            item["score"] = round(
                score,
                4
            )

            results.append(item)

    return results


# =========================================================
# GEMINI ANSWER GENERATION
# =========================================================

def generate_answer(
    question,
    retrieved
):

    if not retrieved:

        return {

            "answer":
                "I couldn't find this information in the uploaded document.",

            "sources": []

        }

    context_parts = []

    for number, item in enumerate(
        retrieved,
        start=1
    ):

        context_parts.append(
            f"""
SOURCE {number}

Document:
{item["filename"]}

Page / Slide:
{item["page"]}

Content:
{item["text"]}
"""
        )

    context = "\n".join(
        context_parts
    )

    prompt = f"""
You are DocuMind AI, a document-grounded RAG assistant.

Answer the user's question ONLY using the document context provided below.

STRICT RULES:

1. Do not use outside knowledge.
2. Do not guess.
3. Do not invent facts.
4. If the answer is not supported by the document context, respond exactly:
"I couldn't find this information in the uploaded document."
5. Keep the answer clear and useful.
6. Preserve important names, dates, numbers and terminology.
7. When appropriate, mention the document name and page/slide.

DOCUMENT CONTEXT:

{context}

USER QUESTION:

{question}

ANSWER:
"""

    last_error = None

    for attempt in range(MAX_RETRIES):

        try:

            response = client.models.generate_content(

                model=GENERATION_MODEL,

                contents=prompt,

                config=types.GenerateContentConfig(

                    temperature=0.1,

                    max_output_tokens=1000

                )

            )

            answer = (

                response.text

                if response.text

                else
                "I couldn't generate an answer from the uploaded document."
            )

            sources = []

            seen = set()

            for item in retrieved:

                key = (
                    item["filename"],
                    item["page"]
                )

                if key in seen:
                    continue

                seen.add(key)

                sources.append({

                    "filename":
                        item["filename"],

                    "page":
                        item["page"],

                    "score":
                        item["score"]

                })

            return {

                "answer":
                    answer.strip(),

                "sources":
                    sources

            }

        except Exception as error:

            last_error = error

            print(
                f"Generation attempt "
                f"{attempt + 1} failed:"
            )

            print(str(error))

            if attempt < MAX_RETRIES - 1:

                import time

                time.sleep(
                    2 ** attempt
                )

    raise last_error


# =========================================================
# DOCUMENT PROCESSING
# =========================================================

def process_document(
    document_id,
    filename,
    pages
):

    try:

        print("\n====================================")

        print("STARTING DOCUMENT INDEXING")

        print(
            f"Document: {filename}"
        )

        print(
            f"Pages: {len(pages)}"
        )

        print("====================================\n")

        new_chunks = create_chunks(
            pages,
            filename,
            document_id
        )

        total_chunks = len(
            new_chunks
        )

        documents[document_id][
            "chunks"
        ] = total_chunks

        print(
            f"Created {total_chunks} chunks."
        )

        def update_progress(
            completed,
            total
        ):

            percent = int(
                (completed / total)
                * 100
            )

            documents[document_id][
                "progress"
            ] = percent

            print(
                f"Progress: {percent}%"
            )

        chunk_count = add_document_to_index(
            new_chunks,
            update_progress
        )

        documents[document_id][
            "chunks"
        ] = chunk_count

        documents[document_id][
            "progress"
        ] = 100

        documents[document_id][
            "status"
        ] = "ready"

        print(
            "\n===================================="
        )

        print("DOCUMENT READY")

        print(
            f"Document: {filename}"
        )

        print(
            f"Chunks: {chunk_count}"
        )

        print(
            "====================================\n"
        )

    except Exception as error:

        documents[document_id][
            "status"
        ] = "error"

        documents[document_id][
            "error"
        ] = str(error)

        print(
            "\n===================================="
        )

        print("DOCUMENT INDEXING ERROR")

        print(str(error))

        print(
            "====================================\n"
        )


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# =========================================================
# UPLOAD
# =========================================================

@app.route(
    "/upload",
    methods=["POST"]
)
def upload():

    if not client:

        return jsonify({

            "success": False,

            "message":
                "GEMINI_API_KEY is not configured."

        }), 500

    if "document" not in request.files:

        return jsonify({

            "success": False,

            "message":
                "No document selected."

        }), 400

    file = request.files[
        "document"
    ]

    if not file.filename:

        return jsonify({

            "success": False,

            "message":
                "Please select a document."

        }), 400

    if not allowed_file(
        file.filename
    ):

        return jsonify({

            "success": False,

            "message":
                "Only PDF, DOCX, TXT and PPTX files are supported."

        }), 400

    filename = secure_filename(
        file.filename
    )

    extension = get_extension(
        filename
    )

    document_id = uuid.uuid4().hex

    unique_filename = (
        f"{document_id}_{filename}"
    )

    file_path = (
        UPLOAD_FOLDER /
        unique_filename
    )

    try:

        file.save(file_path)

        print(
            f"\nUploaded: {filename}"
        )

        pages = extract_document(
            file_path,
            extension
        )

        if not pages:

            file_path.unlink(
                missing_ok=True
            )

            return jsonify({

                "success": False,

                "message":
                    "No readable text was found inside this document."

            }), 400

        total_characters = sum(
            len(page["text"])
            for page in pages
        )

        document_info = {

            "id":
                document_id,

            "filename":
                filename,

            "pages":
                len(pages),

            "characters":
                total_characters,

            "chunks":
                0,

            "progress":
                0,

            "status":
                "processing",

            "error":
                None

        }

        documents[
            document_id
        ] = document_info

        thread = threading.Thread(

            target=process_document,

            args=(
                document_id,
                filename,
                pages
            ),

            daemon=True

        )

        thread.start()

        return jsonify({

            "success":
                True,

            "message":
                "Document uploaded. Indexing started.",

            "document_id":
                document_id,

            "filename":
                filename,

            "pages":
                len(pages),

            "characters":
                total_characters,

            "status":
                "processing"

        })

    except Exception as error:

        print(
            "UPLOAD ERROR:",
            repr(error)
        )

        file_path.unlink(
            missing_ok=True
        )

        return jsonify({

            "success":
                False,

            "message":
                "Document upload failed.",

            "error":
                str(error)

        }), 500


# =========================================================
# DOCUMENT STATUS
# =========================================================

@app.route(
    "/status/<document_id>"
)
def document_status(
    document_id
):

    document = documents.get(
        document_id
    )

    if not document:

        return jsonify({

            "success":
                False,

            "message":
                "Document not found."

        }), 404

    return jsonify({

        "success":
            True,

        "document":
            document

    })


# =========================================================
# DOCUMENT LIST
# =========================================================

@app.route("/documents")
def get_documents():

    return jsonify({

        "success":
            True,

        "documents":
            list(documents.values())

    })


# =========================================================
# CHAT
# =========================================================

@app.route(
    "/chat",
    methods=["POST"]
)
def chat():

    if not client:

        return jsonify({

            "success":
                False,

            "answer":
                "GEMINI_API_KEY is not configured."

        }), 500

    data = request.get_json(
        silent=True
    ) or {}

    question = data.get(
        "message",
        ""
    ).strip()

    if not question:

        return jsonify({

            "success":
                False,

            "answer":
                "Please enter your question."

        }), 400

    ready_documents = [

        document

        for document
        in documents.values()

        if document["status"]
        == "ready"

    ]

    if not ready_documents:

        processing_documents = [

            document

            for document
            in documents.values()

            if document["status"]
            == "processing"

        ]

        if processing_documents:

            return jsonify({

                "success":
                    True,

                "answer":
                    "Your document is still being indexed. Please wait until the status becomes Ready.",

                "sources":
                    []

            })

        return jsonify({

            "success":
                True,

            "answer":
                "Please upload a document first.",

            "sources":
                []

        })

    try:

        retrieved = search_documents(
            question
        )

        result = generate_answer(
            question,
            retrieved
        )

        return jsonify({

            "success":
                True,

            "answer":
                result["answer"],

            "sources":
                result["sources"]

        })

    except Exception as error:

        print(
            "\n========== CHAT ERROR =========="
        )

        print(
            repr(error)
        )

        print(
            "================================\n"
        )

        return jsonify({

            "success":
                False,

            "answer":
                "Unable to process your question.",

            "error":
                str(error)

        }), 500


# =========================================================
# HEALTH
# =========================================================

@app.route("/health")
def health():

    return jsonify({

        "status":
            "ok",

        "service":
            "DocuMind AI",

        "gemini_configured":
            bool(client),

        "generation_model":
            GENERATION_MODEL,

        "embedding_model":
            LOCAL_EMBEDDING_MODEL,

        "embedding_dimension":
            EMBEDDING_DIMENSION,

        "documents":
            len(documents),

        "chunks":
            len(chunks)

    })


# =========================================================
# START SERVER
# =========================================================

if __name__ == "__main__":

    print(
        "\n=========================================="
    )

    print(
        "          DocuMind AI - RAG"
    )

    print(
        "=========================================="
    )

    print(
        "Gemini configured:",
        bool(client)
    )

    print(
        "Generation model:",
        GENERATION_MODEL
    )

    print(
        "Local embedding model:",
        LOCAL_EMBEDDING_MODEL
    )

    print(
        "Embedding dimension:",
        EMBEDDING_DIMENSION
    )

    print(
        "Server:",
        "http://127.0.0.1:5000"
    )

    print(
        "==========================================\n"
    )

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )